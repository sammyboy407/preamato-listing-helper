"""Loads the Master File(s), the Pictures & Measurements file(s), and joins
them.

Only products present in the measurements file are eligible for listing —
that file is the gate, since a listing needs photos. Multiple files of each
type are merged into one combined set, e.g. separate master file exports
per supplier batch, or separate measurements exports per photography batch.
"""
from __future__ import annotations

import csv
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl


@dataclass
class Product:
    sku: str
    master: dict[str, Any]
    measurements: dict[str, Any]

    def m(self, key: str, default=None):
        return self.master.get(key, default)

    def meas(self, key: str, default=None):
        return self.measurements.get(key, default)


def _as_list(paths: str | Path | list) -> list:
    return paths if isinstance(paths, list) else [paths]


def resolve_sheet_name(found: list, wanted: str, path) -> str:
    """Which sheet to read, given the sheets a workbook actually has.

    14.09.26. Uploading the working copy of a sheet instead of the master
    built from it produced a raw Python traceback — "KeyError: 'Worksheet
    Stock Parcel does not exist.'" — which says nothing about what to do
    instead. It happened twice in one afternoon, because the two files are
    indistinguishable in a file picker: the only difference is the tab
    name.

    A workbook with exactly one sheet is not ambiguous, so it is used
    whatever that sheet is called. Anything else names the file, the tab
    it wanted, and the tabs it found.

    Split out as its own function over plain strings deliberately. The
    first test for it built workbooks with openpyxl and died twice on the
    Mac that gates every deploy — no Workbook.active, then no
    Workbook.worksheets — testing that machine's openpyxl rather than this
    decision. The decision is the part that can be wrong.
    """
    if wanted in found:
        return wanted
    if len(found) == 1:
        return found[0]
    raise ValueError(
        f"{Path(path).name} has no sheet called {wanted!r}. Its sheets are: "
        f"{', '.join(repr(n) for n in found)}. The Stock Data File is the one "
        f"with a {wanted!r} tab — check the tab name at the bottom of the "
        f"workbook before uploading."
    )


def load_master_file(path: str | Path, sheet_name: str = "Stock Parcel") -> dict[str, dict[str, Any]]:
    """Returns {SKU: row_dict} for every row in one master file's 'Stock Parcel' sheet."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    # 14.09.26. Uploading the wrong workbook produced a raw Python
    # traceback — "KeyError: 'Worksheet Stock Parcel does not exist.'" —
    # which says nothing about which file to use instead. It happened
    # twice in one afternoon, both times because the working copy of a
    # sheet (tab: "Sheet1") looks identical in a file picker to the master
    # built from it (tab: "Stock Parcel").
    #
    # A workbook with exactly one sheet is not ambiguous, so use it. Any
    # other mismatch names the file, the tab it wanted, and the tabs it
    # actually found.
    found = list(getattr(wb, "sheetnames", None) or wb.get_sheet_names())
    sheet_name = resolve_sheet_name(found, sheet_name, path)
    ws = wb[sheet_name]
    rows = ws.iter_rows(values_only=True)
    headers = [h.strip() if isinstance(h, str) else h for h in next(rows)]
    out = {}
    for row in rows:
        record = dict(zip(headers, row))
        sku = record.get("SKU")
        if sku:
            out[str(sku).strip()] = record
    wb.close()
    return out


def load_master_files(paths: str | Path | list, sheet_name: str = "Stock Parcel") -> dict[str, dict[str, Any]]:
    """Merges multiple master files. A SKU appearing in more than one file
    uses the last file's row (files are merged in the order given) — a
    warning is printed so an accidental duplicate export doesn't silently
    shadow real data."""
    merged: dict[str, dict[str, Any]] = {}
    seen_in: dict[str, str] = {}
    for path in _as_list(paths):
        for sku, record in load_master_file(path, sheet_name).items():
            if sku in merged and seen_in[sku] != str(path):
                print(f"WARNING: SKU {sku!r} appears in multiple master files — "
                      f"using the row from {path} (last one wins).")
            merged[sku] = record
            seen_in[sku] = str(path)
    return merged


def load_measurements(path: str | Path) -> dict[str, dict[str, Any]]:
    """Returns {SKU: row_dict} keyed on one measurements file's 'Name' column
    (which holds the product SKU — the file's own 'SKU' column is actually
    the barcode)."""
    out = {}
    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        reader.fieldnames = [h.strip() if h else h for h in (reader.fieldnames or [])]
        for record in reader:
            sku = (record.get("Name") or "").strip()
            if sku:
                out[sku] = record
    return out


def load_measurements_files(paths: str | Path | list) -> dict[str, dict[str, Any]]:
    """Merges multiple measurements files, same last-wins-with-warning rule
    as load_master_files."""
    merged: dict[str, dict[str, Any]] = {}
    seen_in: dict[str, str] = {}
    for path in _as_list(paths):
        for sku, record in load_measurements(path).items():
            if sku in merged and seen_in[sku] != str(path):
                print(f"WARNING: SKU {sku!r} appears in multiple measurements files — "
                      f"using the row from {path} (last one wins).")
            merged[sku] = record
            seen_in[sku] = str(path)
    return merged


def split_image_urls(raw: str | None) -> list[str]:
    if not raw:
        return []
    return [u.strip() for u in raw.replace("\r", "\n").split("\n") if u.strip()]


def load_products(
    master_path: str | Path | list, measurements_path: str | Path | list
) -> tuple[list[Product], list[str]]:
    """Returns (products, unmatched_skus). unmatched_skus is every SKU present
    in a measurements file with no matching row in any master file — most
    often a stock file that was never uploaded alongside the measurements
    for that batch (BRK02, 24.09.26: the measurements loaded fine, the
    matching Stock Data File wasn't on the run, and every one of its SKUs
    vanished with nothing but a line in a log nobody was reading).

    Previously only a print() here — invisible in the Streamlit UI, which
    reads from pipeline.run()'s return values, not stdout. Now returned so
    the caller can put it wherever failed/uncovered/held_back already show
    up. The print stays too, for the CLI running with nothing else watching."""
    master = load_master_files(master_path)
    measurements = load_measurements_files(measurements_path)

    products = []
    unmatched = []
    for sku, meas_row in measurements.items():
        master_row = master.get(sku)
        if master_row is None:
            unmatched.append(sku)
            continue
        products.append(Product(sku=sku, master=master_row, measurements=meas_row))

    if unmatched:
        print(f"WARNING: {len(unmatched)} SKU(s) in measurements file(s) have no match in "
              f"the master file(s) and will be skipped: {unmatched}")

    return products, unmatched
